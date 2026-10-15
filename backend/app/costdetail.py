"""费用明细：每项成本科目按 时间 / 市场 / 区块 / 队伍 汇总与导出，供专业部室做单项成本对比分析。

数据全部来自经济事项（只取“有效”，冲销与被冲销互相抵消），与经营总览同源，金额不另算。
权限沿用经营报表：部室只看归口给本部室的科目、看不到收入；经营总览（企划财务部）看全部成本科目。
口径：市场、区块取单井基础信息；队伍取事项发生时的施工队伍（事项上没有队伍时按该井当期施工队伍补）。
进尺：已审核日报的进尺合计（米）；钻机台月：施工时段落在各月的天数 ÷ 当月天数（整月施工为 1 台月）。
米成本 = 金额 ÷ 进尺（元/米）；单位成本 = 金额 ÷ 钻机台月（元/台月）。分母取所选期间内有经济事项的全部井·月，不随所选科目变化，
所以单项科目的米成本、单位成本可以直接横向对比。
米成本的金额只统计有进尺的井·月（模板导入的历史数据没有进尺，不进米成本的分子，避免虚高）；单位成本统计全部井·月。
"""
import calendar
from datetime import date
import re
from io import BytesIO

import openpyxl
from fastapi import HTTPException
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import analysis, auth
from .db import now

DIMS = {'period': '时间', 'market': '市场', 'block': '区块', 'team': '队伍'}
NONE_LABEL = '未归属'
PERIOD_RE = re.compile(r'^\d{4}-(0[1-9]|1[0-2])$')


def _bad(msg: str):
    return HTTPException(400, msg)


def _check_range(start: str | None, end: str | None) -> tuple[str | None, str | None]:
    start, end = start or None, end or None
    for p in (start, end):
        if p and not PERIOD_RE.match(p):
            raise _bad('期间格式应为 YYYY-MM')
    if start and end and start > end:
        raise _bad('起始期间不能晚于截止期间')
    return start, end


def _scope_items(conn, user: dict, codes: list[str] | None) -> list[dict]:
    """当前账号可见的成本科目（树序）。codes 非空时只保留所选科目及其下级。"""
    auth.require(user, 'finance', 'dept')
    items = [dict(r) for r in conn.execute("SELECT * FROM cost_item WHERE direction='成本' ORDER BY sort_no")]
    children: dict[str, list[str]] = {}
    for it in items:
        if it['parent_code']:
            children.setdefault(it['parent_code'], []).append(it['item_code'])
    visible: dict[str, bool] = {}
    for it in reversed(items):   # 子项排在父项之后，倒序即自底向上
        c = it['item_code']
        visible[c] = auth.item_visible(user, it) if it['is_leaf'] else any(visible.get(k) for k in children.get(c, []))
    keep = {c for c, v in visible.items() if v}
    if codes:
        unknown = [c for c in codes if c not in keep]
        if unknown:
            raise _bad('科目不存在或不在可见范围内：' + '、'.join(unknown))
        picked: set[str] = set()
        stack = list(codes)
        while stack:
            c = stack.pop()
            if c in keep and c not in picked:
                picked.add(c)
                stack.extend(children.get(c, []))
        keep = picked
    return [{'code': it['item_code'], 'name': it['item_name'], 'level': it['level'], 'is_leaf': bool(it['is_leaf']),
             'parent': it['parent_code'], 'dept': it['owner_dept'] or '', 'behavior': it['cost_behavior'] or ''}
            for it in items if it['item_code'] in keep]


def _events(conn, items: list[dict], start: str | None, end: str | None) -> list[dict]:
    leaves = [i['code'] for i in items if i['is_leaf']]
    if not leaves:
        return []
    sql = ("SELECT e.event_id, e.event_date, e.period, e.well_id, e.item_code, e.amount, e.quantity, e.unit_price, e.origin_type, "
           "e.team_id, w.market, w.block, t.team_no, t.name team_name, f.unit "
           "FROM economic_event e JOIN well w ON w.well_id=e.well_id LEFT JOIN team t ON t.team_id=e.team_id "
           "LEFT JOIN business_fact f ON f.fact_id=e.fact_id "
           f"WHERE e.status='有效' AND e.well_id NOT IN (SELECT well_id FROM well_deleted) AND e.item_code IN ({','.join('?' * len(leaves))})")
    args: list = list(leaves)
    if start:
        sql += ' AND e.period>=?'
        args.append(start)
    if end:
        sql += ' AND e.period<=?'
        args.append(end)
    out = []
    for r in conn.execute(sql + ' ORDER BY e.item_code, e.period, e.well_id, e.event_id', args):
        d = dict(r)
        if not d['team_id']:   # 历史事项可能没有队伍：按该井当期施工队伍补
            seg = analysis._well_segment(conn, d['well_id'], d['period'])
            if seg:
                d['team_no'], d['team_name'] = seg['team_no'], seg['team_name']
        d['amount'] = int(round(d['amount']))   # 每笔事项入库时已取整到元
        d['market'] = d['market'] or NONE_LABEL
        d['block'] = d['block'] or NONE_LABEL
        d['team_name'] = d['team_name'] or NONE_LABEL
        d['team_no'] = d['team_no'] or ''
        out.append(d)
    return out


def _denoms(conn, start: str | None, end: str | None) -> list[dict]:
    """进尺（米）与钻机台月：按 井·月·队伍 汇总，供各维度求米成本和单位成本。"""
    sql = ("SELECT DISTINCT e.period, e.well_id, w.market, w.block FROM economic_event e JOIN well w ON w.well_id=e.well_id "
           "WHERE e.status='有效' AND e.well_id NOT IN (SELECT well_id FROM well_deleted)")
    args: list = []
    if start:
        sql += ' AND e.period>=?'
        args.append(start)
    if end:
        sql += ' AND e.period<=?'
        args.append(end)
    out = []
    for r in conn.execute(sql, args).fetchall():
        period, wid = r['period'], r['well_id']
        y, m = int(period[:4]), int(period[5:])
        days_in = calendar.monthrange(y, m)[1]
        first, last = f'{period}-01', f'{period}-{days_in:02d}'
        teams: dict[str, dict] = {}
        for sg in conn.execute('SELECT p.start_at, p.end_at, t.team_no, t.name FROM team_well_period p JOIN team t ON t.team_id=p.team_id '
                               'WHERE p.well_id=?', (wid,)):
            a, b = max(sg['start_at'][:10], first), min((sg['end_at'] or '9999-12-31')[:10], last)
            if a <= b:
                d = (date.fromisoformat(b) - date.fromisoformat(a)).days + 1
                t = teams.setdefault(sg['name'], {'team_no': sg['team_no'], 'footage': 0.0, 'days': 0})
                t['days'] += d
        for f in conn.execute("SELECT t.name, t.team_no, SUM(f.quantity) q FROM business_fact f JOIN team t ON t.team_id=f.team_id "
                              "WHERE f.well_id=? AND f.fact_type='进尺' AND f.status='已审核' AND substr(f.fact_date,1,7)=? "
                              'GROUP BY f.team_id', (wid, period)):
            t = teams.setdefault(f['name'], {'team_no': f['team_no'], 'footage': 0.0, 'days': 0})
            t['footage'] += f['q'] or 0
        if not teams:
            teams[NONE_LABEL] = {'team_no': '', 'footage': 0.0, 'days': 0}
        for name, t in teams.items():
            out.append({'period': period, 'well_id': wid, 'market': r['market'] or NONE_LABEL, 'block': r['block'] or NONE_LABEL,
                        'team_name': name, 'team_no': t['team_no'], 'footage': t['footage'], 'rig_months': t['days'] / days_in})
    return out


def _per(amount: float, base: float) -> float | None:
    return round(amount / base, 2) if base else None


def _dim_key(ev: dict, by: str) -> str:
    return {'period': ev['period'], 'market': ev['market'], 'block': ev['block'], 'team': ev['team_name']}[by]


def options(conn, user: dict) -> dict:
    """页面筛选项：可见科目树、有数据的期间，以及各维度实际出现的对象。"""
    items = _scope_items(conn, user, None)
    evs = _events(conn, items, None, None)
    return {
        'items': items,
        'periods': sorted({e['period'] for e in evs}),
        'market': sorted({e['market'] for e in evs}),
        'block': sorted({e['block'] for e in evs}),
        'team': [t for _, t in sorted({(e['team_no'], e['team_name']) for e in evs})],
        'scope': user['scope_label'],
    }


def matrix(conn, user: dict, by: str, codes: list[str] | None = None, start: str | None = None, end: str | None = None) -> dict:
    """行 = 成本科目（含上级小计），列 = 所选维度的各个对象，格 = 金额（元）。"""
    if by not in DIMS:
        raise _bad('维度应为 period / market / block / team 之一')
    start, end = _check_range(start, end)
    items = _scope_items(conn, user, codes)
    evs = _events(conn, items, start, end)
    index = {i['code']: i for i in items}

    order: dict[str, tuple] = {}   # 列的排序键：队伍按队号，其余按名称
    wells: dict[str, set] = {}
    cell: dict[tuple[str, str], int] = {}
    cellf: dict[tuple[str, str], int] = {}   # 只含有进尺的井·月，作米成本的分子
    totals: dict[str, int] = {}
    totalsf: dict[str, int] = {}
    denoms = _denoms(conn, start, end)
    has_foot = {(d['well_id'], d['period']) for d in denoms if d['footage'] > 0}
    for e in evs:
        k = _dim_key(e, by)
        withf = (e['well_id'], e['period']) in has_foot
        if withf:
            totalsf[k] = totalsf.get(k, 0) + e['amount']
        order.setdefault(k, (e['team_no'], k) if by == 'team' else (k,))
        wells.setdefault(k, set()).add(e['well_id'])
        totals[k] = totals.get(k, 0) + e['amount']
        code = e['item_code']
        while code and code in index:   # 叶子金额逐级累加到上级
            cell[(code, k)] = cell.get((code, k), 0) + e['amount']
            if withf:
                cellf[(code, k)] = cellf.get((code, k), 0) + e['amount']
            code = index[code]['parent']
    cols = sorted(order, key=order.__getitem__)
    foot: dict[str, float] = {k: 0.0 for k in cols}
    rig: dict[str, float] = {k: 0.0 for k in cols}
    for d in denoms:
        k = _dim_key(d, by)
        if k in foot:
            foot[k] += d['footage']
            rig[k] += d['rig_months']
    rows = []
    for it in items:
        vals = {k: cell.get((it['code'], k), 0) for k in cols}
        tot = sum(vals.values())
        valsf = {k: cellf.get((it['code'], k), 0) for k in cols}
        rows.append({**it, 'values': vals, 'total': tot,
                     'per_meter': {k: _per(v, foot[k]) for k, v in valsf.items()}, 'per_rig': {k: _per(v, rig[k]) for k, v in vals.items()},
                     'per_meter_total': _per(sum(valsf.values()), sum(foot.values())), 'per_rig_total': _per(tot, sum(rig.values()))})
    grand = sum(totals.values())
    columns = [{'key': k, 'label': k, 'wells': len(wells[k]), 'total': totals[k], 'footage': round(foot[k], 2),
                'rig_months': round(rig[k], 2), 'cost_with_footage': totalsf.get(k, 0),
                'per_meter': _per(totalsf.get(k, 0), foot[k]), 'per_rig': _per(totals[k], rig[k])} for k in cols]
    return {'by': by, 'label': DIMS[by], 'start': start, 'end': end, 'scope': user['scope_label'],
            'columns': columns, 'rows': rows, 'grand_total': grand, 'event_count': len(evs),
            'footage': round(sum(foot.values()), 2), 'rig_months': round(sum(rig.values()), 2),
            'cost_with_footage': sum(totalsf.values()),
            'per_meter': _per(sum(totalsf.values()), sum(foot.values())), 'per_rig': _per(grand, sum(rig.values()))}


def detail(conn, user: dict, codes: list[str] | None = None, start: str | None = None, end: str | None = None) -> list[dict]:
    start, end = _check_range(start, end)
    items = _scope_items(conn, user, codes)
    names = {i['code']: i for i in items}
    return [{**e, 'item_name': names[e['item_code']]['name'], 'dept': names[e['item_code']]['dept']}
            for e in _events(conn, items, start, end)]


# ---------- Excel ----------

HEAD_FILL = PatternFill('solid', fgColor='E8E8ED')
SUB_FILL = PatternFill('solid', fgColor='F5F5F7')
NOTE = ('口径：经济事项有效金额合计（冲销已抵消）；金额单位：元；市场、区块取单井基础信息；队伍取事项发生时的施工队伍。'
        '进尺为已审核日报进尺合计（米）；钻机台月 = 施工天数 ÷ 当月天数；米成本 = 有进尺井·月的金额 ÷ 进尺（元/米）；单位成本 = 金额 ÷ 钻机台月（元/台月）。')
METRICS = {'amount': ('金额（元）', None), 'per_meter': ('米成本（元/米）', 'per_meter'), 'per_rig': ('单位成本（元/台月）', 'per_rig')}


def _range_text(start: str | None, end: str | None) -> str:
    return f'{start} 至 {end}' if start and end else (f'{start} 起' if start else (f'截至 {end}' if end else '全部期间'))


def _write_matrix(ws, m: dict, metric: str = 'amount') -> None:
    mlabel = METRICS[metric][0]
    money_fmt = '#,##0;[Red]-#,##0;"-"' if metric == 'amount' else '#,##0.00;[Red]-#,##0.00;"-"'
    ws.append([f'费用明细 · 按{m["label"]} · {mlabel} · {m["scope"]} · 期间：{_range_text(m["start"], m["end"])}'])
    ws.append([NOTE + f' 导出时间：{now()[:16]}'])
    ws['A1'].font = Font(bold=True, size=13)
    ws['A2'].font = Font(color='6E6E73', size=10)
    ncol = len(m['columns'])
    ws.append(['科目编码', '科目名称', '归口部室', '成本性质'] + [c['label'] for c in m['columns']] + ['合计'])
    ws.append(['', '涉及井数', '', ''] + [c['wells'] for c in m['columns']] + [''])
    ws.append(['', '进尺（米）', '', ''] + [c['footage'] for c in m['columns']] + [m['footage']])
    ws.append(['', '钻机台月', '', ''] + [c['rig_months'] for c in m['columns']] + [m['rig_months']])
    for c in ws[3]:
        c.font, c.fill = Font(bold=True), HEAD_FILL
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for row in (4, 5, 6):
        for c in ws[row]:
            c.fill, c.font = SUB_FILL, Font(color='6E6E73')
    for col in range(5, 5 + ncol + 1):
        ws.cell(5, col).number_format = '#,##0.0'
        ws.cell(6, col).number_format = '#,##0.00'

    def cellval(r: dict, c: dict | None):
        if metric == 'amount':
            return r['total'] if c is None else r['values'][c['key']]
        key = METRICS[metric][1]
        return r[key + '_total'] if c is None else r[key][c['key']]

    for r in m['rows']:
        ws.append([r['code'], r['name'], r['dept'], r['behavior']] + [cellval(r, c) for c in m['columns']] + [cellval(r, None)])
        row = ws.max_row
        ws.cell(row, 2).alignment = Alignment(indent=max(r['level'] - 1, 0))
        for col in range(5, 5 + ncol + 1):
            ws.cell(row, col).number_format = money_fmt
        if not r['is_leaf']:
            for col in range(1, 5 + ncol + 1):
                ws.cell(row, col).font = Font(bold=True)
    if m['rows']:
        if metric == 'amount':
            tot = [c['total'] for c in m['columns']] + [m['grand_total']]
        else:
            key = METRICS[metric][1]
            tot = [c[key] for c in m['columns']] + [m[key]]
        ws.append(['', '合计（叶子科目）', '', ''] + tot)
        for col in range(1, 5 + ncol + 1):
            cell = ws.cell(ws.max_row, col)
            cell.font, cell.fill = Font(bold=True), HEAD_FILL
            if col >= 5:
                cell.number_format = money_fmt
        if metric == 'amount':   # 金额表底部附总成本口径的两个单位指标
            for label, key in (('米成本（元/米）', 'per_meter'), ('单位成本（元/台月）', 'per_rig')):
                ws.append(['', label, '', ''] + [c[key] for c in m['columns']] + [m[key]])
                for col in range(1, 5 + ncol + 1):
                    cell = ws.cell(ws.max_row, col)
                    cell.font, cell.fill = Font(bold=True, color='6E6E73'), SUB_FILL
                    if col >= 5:
                        cell.number_format = '#,##0.00;[Red]-#,##0.00;"-"'
    ws.column_dimensions['A'].width = 14
    ws.column_dimensions['B'].width = 30
    ws.column_dimensions['C'].width = 10
    ws.column_dimensions['D'].width = 9
    for i in range(5, 5 + ncol + 1):
        ws.column_dimensions[get_column_letter(i)].width = 16
    ws.freeze_panes = 'E7'


def workbook(conn, user: dict, codes: list[str] | None, start: str | None, end: str | None, bys: list[str], with_detail: bool = True,
             with_unit: bool = True) -> bytes:
    bys = [b for b in dict.fromkeys(bys) if b in DIMS] or list(DIMS)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for b in bys:
        m = matrix(conn, user, b, codes, start, end)
        _write_matrix(wb.create_sheet('按' + DIMS[b]), m)
        if with_unit:   # 米成本、单位成本各一张同结构的表，便于单项科目横向对比
            _write_matrix(wb.create_sheet(f'按{DIMS[b]}·米成本'), m, 'per_meter')
            _write_matrix(wb.create_sheet(f'按{DIMS[b]}·台月成本'), m, 'per_rig')
    if with_detail:
        ws = wb.create_sheet('事项明细')
        ws.append(['期间', '日期', '井号', '市场', '区块', '队号', '队伍', '科目编码', '科目名称', '归口部室', '数量', '单位', '单价（元）',
                   '金额（元）', '来源', '事项编号'])
        for c in ws[1]:
            c.font, c.fill = Font(bold=True), HEAD_FILL
        for e in detail(conn, user, codes, start, end):
            ws.append([e['period'], e['event_date'], e['well_id'], e['market'], e['block'], e['team_no'], e['team_name'], e['item_code'],
                       e['item_name'], e['dept'], e['quantity'], e['unit'] or '', e['unit_price'], e['amount'], e['origin_type'], e['event_id']])
            ws.cell(ws.max_row, 14).number_format = '#,##0;[Red]-#,##0'
        for i, w in enumerate([9, 11, 9, 10, 10, 8, 14, 13, 28, 10, 11, 8, 11, 13, 10, 9], start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
