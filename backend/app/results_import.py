"""单井经营结果模板导入：按《单井经营情况导出模板》列转行，以“历史导入”写入经济事项。

一个工作表里横向放多口井：A 列为项目，B 列起每一列是一口井一个月（第 1 行井号、年份、月份等表头在每列各填一次），
行序与模板一致；只有一列 B 的旧格式（一个工作表一口井一个月）同样支持。先预览、后确认。
- 金额按末级项目导入，汇总行由系统重算；汇总行与明细不一致时提示。
- 同一井同一月再次导入：逐个科目比对，金额有变化的先冲销旧记录再导入新金额（只冲不改），没变的科目保留，全部相同记为无变化。
- 支持一次选多个文件，每个文件可含多个工作表、每个工作表横向放多口井；同一井同一月不能重复出现。
- 已手工删除的井：历史数据井导入时自动恢复；已由基层队登记过的井须由基层队重新登记才能恢复。
- 该井该月已有规则计算的事项时不允许导入，避免与系统计算重复。
- 新井按表头建档、按队号挂施工时段；已有井以单井基础信息为准，表头不一致只提示。
- 未经基层钻井队登记基础信息的井在经营总览备注“历史数据井”；已登记的井含历史导入金额时备注“历史数据导入”。
"""
import copy
import re
from io import BytesIO

import openpyxl
from openpyxl.utils import get_column_letter

from . import engine
from .db import log, money, now
from .excel import TEMPLATE, whole_yuan_format
from .wellinfo import SECTIONS

HEAD_ROWS = 18
TEMPLATE_WELLS = 10   # 空白模板预留的井列数（B 列起），列数不限，多填几列即可
SUMMARY_LABELS = ['经营数据', '总收入（元）', '总成本（元）', '（一）固定成本（元）', '（二）变动成本（元）', '毛益额（元）', '毛益率（%）',
                  '净利润（元）', '利润率（%）']


class ImportError_(Exception):
    pass


def template_labels() -> list[str]:
    ws = openpyxl.load_workbook(TEMPLATE).active
    return [str(ws.cell(r, 1).value or '').strip() for r in range(1, ws.max_row + 1)]


def blank_template() -> bytes:
    """空白批量模板：A 列项目，B–K 列各放一口井一个月（样式沿用单井模板的 B 列）。"""
    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb.active
    whole_yuan_format(ws)
    width = ws.column_dimensions['B'].width
    for r in range(1, ws.max_row + 1):
        src = ws.cell(r, 2)
        src.value = None
        for col in range(3, 2 + TEMPLATE_WELLS):
            dst = ws.cell(r, col)
            dst._style = copy.copy(src._style)
    for col in range(2, 2 + TEMPLATE_WELLS):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.title = '批量导入'
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _num(v):
    if v is None or v == '':
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(',', '').replace('，', '')
    try:
        return float(s)
    except ValueError:
        return None


def _s(v) -> str:
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def parse(data: bytes) -> list[dict]:
    """每个工作表 A 列为项目，B 列起每个填了井号的列是一口井一个月；返回按列拆开的待校验清单。"""
    try:
        wb = openpyxl.load_workbook(BytesIO(data), data_only=True)
    except Exception as e:  # noqa: BLE001
        raise ImportError_(f'无法读取文件：{e}') from e
    sheets = []
    for ws in wb.worksheets:
        labels = [str(ws.cell(r, 1).value or '').strip() for r in range(1, ws.max_row + 1)]
        if not labels or labels[0] != '井号':
            continue
        cols = [c for c in range(2, ws.max_column + 1) if _s(ws.cell(1, c).value)
                or any(ws.cell(r, c).value not in (None, '') for r in range(2, ws.max_row + 1))]
        for c in cols:
            name = ws.title if cols == [2] else f'{ws.title} · {get_column_letter(c)}列'
            sheets.append({'sheet': name, 'col': c, 'cells': [(labels[r - 1], ws.cell(r, c).value) for r in range(1, ws.max_row + 1)]})
    if not sheets:
        raise ImportError_('没有找到《单井经营情况导出模板》格式的工作表（A1 应为“井号”，井从 B 列起横向排列）')
    return sheets


def _cost_items(conn) -> list[dict]:
    return [dict(r) for r in conn.execute("SELECT * FROM cost_item WHERE direction='成本' ORDER BY sort_no")]


def check(conn, sheet: dict, labels: list[str], items: list[dict], seen: set) -> dict:
    file = sheet.get('file', '')
    cells = sheet['cells']
    blocks, warns = [], []
    if len(cells) < len(labels):
        blocks.append(f'行数 {len(cells)} 少于模板的 {len(labels)} 行')
    for i, lab in enumerate(labels):
        if i < len(cells) and cells[i][0] != lab:
            blocks.append(f'第 {i + 1} 行项目“{cells[i][0]}”与模板“{lab}”不一致')
            if len(blocks) > 5:
                break
    head = {k: v for k, v in cells[:HEAD_ROWS]}
    wid = _s(head.get('井号'))
    y, m = _s(head.get('年份')), _s(head.get('月份'))
    period = f'{int(y):04d}-{int(m):02d}' if y.isdigit() and m.isdigit() and 1 <= int(m) <= 12 else ''
    if not wid:
        blocks.append('井号为空')
    if not period:
        blocks.append('年份、月份不正确')
    if (wid, period) in seen:
        blocks.append('文件内同一井同一月重复')
    seen.add((wid, period))

    amounts, bad = {}, []
    for n, it in enumerate(items):
        row = HEAD_ROWS + n
        if row >= len(cells):
            break
        v = _num(cells[row][1])
        if v is None:
            bad.append(f'第 {row + 1} 行')
            v = 0.0
        amounts[it['item_code']] = v
    if bad:
        blocks.append('金额不是数字：' + '、'.join(bad[:5]) + (' 等' if len(bad) > 5 else ''))
    revenue = _num(head.get('形象收入（元）'))
    if revenue is None:
        blocks.append('形象收入不是数字')
        revenue = 0.0

    # 汇总行与明细比对：按末级明细导入，汇总由系统重算
    children: dict[str, list[str]] = {}
    for it in items:
        if it['parent_code']:
            children.setdefault(it['parent_code'], []).append(it['item_code'])
    calc: dict[str, float] = {}
    for it in reversed(items):
        c = it['item_code']
        calc[c] = amounts.get(c, 0.0) if it['is_leaf'] else sum(calc[k] for k in children.get(c, []))
    for it in items:
        if not it['is_leaf'] and abs(calc[it['item_code']] - amounts.get(it['item_code'], 0.0)) > 1:
            warns.append(f'{it["item_name"]} 合计 {amounts.get(it["item_code"], 0):,.2f} 与明细合计 {calc[it["item_code"]]:,.2f} 不一致，按明细导入')
    for rr in conn.execute("SELECT DISTINCT r.item_code, i.item_name, r.unit_price FROM economic_rule r JOIN cost_item i "
                           "ON i.item_code=r.item_code WHERE r.status='已发布' AND r.calc_order=3"):
        if revenue and money(amounts.get(rr['item_code'], 0.0)) == 0:
            warns.append(f'{rr["item_name"]} 为 0，导入后按已发布规则计提（形象收入 × {rr["unit_price"] * 100:g}%）')
    total = calc.get('C1', 0) + calc.get('C2', 0)
    tc = _num(head.get('总成本（元）'))
    if tc is not None and abs(tc - total) > 1:
        warns.append(f'总成本 {tc:,.2f} 与明细合计 {total:,.2f} 不一致，按明细导入')

    team_no = _s(head.get('队号'))
    team = None
    if team_no:
        team = conn.execute('SELECT * FROM team WHERE team_no=? OR name=?', (team_no, team_no)).fetchone()
    well = conn.execute('SELECT * FROM well WHERE well_id=?', (wid,)).fetchone() if wid else None
    if well:
        diffs = [lab for lab, col in (('市场分布', 'market'), ('区块', 'block'), ('井别', 'well_category'), ('井型', 'well_type'))
                 if _s(head.get(lab)) and _s(head.get(lab)) != (well[col] or '')]
        if diffs:
            warns.append(f'{"、".join(diffs)}与系统中的单井信息不一致，以系统为准')
    elif wid and not team_no:
        blocks.append('新井必须填写队号')
    deleted = bool(wid) and bool(conn.execute('SELECT 1 FROM well_deleted WHERE well_id=?', (wid,)).fetchone())
    if deleted:
        registered = conn.execute("SELECT 1 FROM well_info_report WHERE well_id=? AND status IN ('已提交', '已被更新') LIMIT 1", (wid,)).fetchone()
        if registered:
            blocks.append(f'{wid} 已被删除，且由基层钻井队登记过；请由基层队重新登记基础信息后再导入')
        else:
            warns.append(f'{wid} 已被删除，导入后将恢复')
    if wid and period:
        engine_n = conn.execute(
            "SELECT COUNT(*) n FROM economic_event WHERE well_id=? AND period=? AND status='有效' AND origin_type IN ('规则计算','比例计提') "
            "AND item_code IN (SELECT item_code FROM cost_item WHERE direction='成本' OR item_code='R02')", (wid, period)).fetchone()['n']
        imported_n = conn.execute("SELECT COUNT(*) n FROM economic_event WHERE well_id=? AND period=? AND origin_type='历史导入' "
                                  "AND status='有效'", (wid, period)).fetchone()['n']
        if engine_n and not imported_n:
            blocks.append(f'{wid} {period} 已有规则计算的经济事项，历史数据只能导入系统上线前的期间')

    new = {c: money(v) for c, v in amounts.items() if _leaf(items, c) and money(v) != 0}   # 入库按整数元
    if money(revenue) != 0:
        new['R02'] = money(revenue)
    cost_in = sum(v for c, v in new.items() if c != 'R02')   # 预览合计 = 入库后各末级科目之和
    old = {}
    if wid and period:
        for r in conn.execute("SELECT item_code, SUM(amount) a FROM economic_event WHERE well_id=? AND period=? AND origin_type='历史导入' "
                              "AND status='有效' GROUP BY item_code", (wid, period)):
            old[r['item_code']] = money(r['a'])
    names = {it['item_code']: it['item_name'] for it in items}
    names['R02'] = '形象收入'
    changes = [{'item': names.get(c, c), 'old': old.get(c, 0), 'new': new.get(c, 0)}
               for c in sorted(set(old) | set(new)) if old.get(c, 0) != new.get(c, 0)] if old else []
    if blocks:
        action = '阻断'
    elif old and old == new:
        action = '无变化'
    elif old:
        action = '覆盖'
    else:
        action = '导入'
    return {'file': file, 'sheet': sheet['sheet'], 'well_id': wid, 'period': period, 'team_no': team_no, 'new_well': bool(wid) and not well,
            'new_team': bool(team_no) and not team, 'revenue': new.get('R02', 0), 'total_cost': cost_in,
            'net_profit': new.get('R02', 0) - cost_in, 'action': action, 'blocks': blocks, 'warns': warns,
            'changes': changes, '_head': head, '_new': new, '_old': old, '_team': dict(team) if team else None}


def _leaf(items, code):
    return any(it['item_code'] == code and it['is_leaf'] for it in items)


def run(conn, files: list[tuple[str, bytes]], user: dict, commit: bool) -> dict:
    """files：[(文件名, 内容)]，可一次导入多个文件、每个文件多个工作表。"""
    labels = template_labels()
    items = _cost_items(conn)
    if len(items) + HEAD_ROWS + len(SUMMARY_LABELS) != len(labels):
        raise ImportError_('系统成本项目树与导出模板行数不一致，请先同步模板')
    seen: set = set()
    results = []
    for filename, data in files:
        try:
            sheets = parse(data)
        except ImportError_ as e:
            if len(files) == 1:
                raise
            results.append({'file': filename, 'sheet': '—', 'well_id': '', 'period': '', 'team_no': '', 'new_well': False,
                            'new_team': False, 'revenue': 0, 'total_cost': 0, 'net_profit': 0, 'action': '阻断',
                            'blocks': [f'{filename}：{e}'], 'warns': [], 'changes': []})
            continue
        for sh in sheets:
            sh['file'] = filename
            results.append(check(conn, sh, labels, items, seen))
    if commit:
        batch_id = conn.execute(
            "INSERT INTO calc_batch (trigger, scope, engine_version, started_at, operator) VALUES ('模板导入', ?, 'import-1.0', ?, ?)",
            ('、'.join(f for f, _ in files)[:200], now(), user['who'])).lastrowid
        n = 0
        for r in results:
            if r['action'] not in ('导入', '覆盖'):
                continue
            n += _write(conn, r, r['file'], batch_id, user)
        conn.execute('UPDATE calc_batch SET finished_at=?, event_count=? WHERE batch_id=?', (now(), n, batch_id))
    count: dict[str, int] = {}
    for r in results:
        count[r['action']] = count.get(r['action'], 0) + 1
    return {'count': count, 'committed': commit,
            'sheets': [{k: v for k, v in r.items() if not k.startswith('_')} for r in results]}


def _write(conn, r: dict, filename: str, batch_id: int, user: dict) -> int:
    """写入一张工作表，返回新写入的事项数；覆盖时只冲销并重写金额有变化的科目。"""
    head, wid, period = r['_head'], r['well_id'], r['period']
    team = r['_team']
    if not team and r['team_no']:   # 同一批里前面的井可能已经建了这支队伍，写入时重新查一次
        team = conn.execute('SELECT * FROM team WHERE team_no=? OR name=?', (r['team_no'], r['team_no'])).fetchone()
    if not team and not r['team_no']:
        tid = None   # 已有井、表中未填队号：沿用系统中的施工时段
    elif not team:
        tid = 'T' + re.sub(r'\s+', '', r['team_no'])
        conn.execute('INSERT INTO team (team_id, team_no, name, rig_type, unit, captain) VALUES (?,?,?,?, NULL, ?)',
                     (tid, r['team_no'], f"{r['team_no']} 钻井队", _s(head.get('钻机类型')) or None, _s(head.get('队长')) or None))
    else:
        tid = team['team_id']
    phase = _s(head.get('开次'))
    if phase.isdigit() and 1 <= int(phase) <= 5:
        phase = SECTIONS[int(phase) - 1]
    depth = _num(head.get('实际井深（米）')) or 0
    contract = money(_num(head.get('合同收入（元）')) or 0)
    conn.execute('DELETE FROM well_deleted WHERE well_id=?', (wid,))   # 已删除的历史数据井导入即恢复
    if not conn.execute('SELECT 1 FROM well WHERE well_id=?', (wid,)).fetchone():   # 同一口井的多个月份在同一批里，只建一次
        conn.execute("INSERT INTO well VALUES (?,?,?,?,?,?,?,?,?,?, '完井')",
                     (wid, _s(head.get('市场分布')) or '—', _s(head.get('区块')) or '—', _s(head.get('井别')) or '—',
                      _s(head.get('井型')) or '—', SECTIONS.index(phase) + 1 if phase in SECTIONS else 1,
                      _num(head.get('设计井深（米）')) or 0, depth, _s(head.get('甲方单位')) or None, contract))
    else:
        conn.execute('UPDATE well SET actual_depth=MAX(actual_depth, ?), contract_amount=CASE WHEN contract_amount=0 THEN ? '
                     'ELSE contract_amount END WHERE well_id=?', (depth, contract, wid))
    start, end = period + '-01', engine.period_end(period)
    if tid and not conn.execute('SELECT 1 FROM team_well_period WHERE well_id=? AND start_at<=? AND (end_at IS NULL OR end_at>=?)',
                        (wid, end, start)).fetchone():
        conn.execute('INSERT INTO team_well_period (team_id, well_id, phase, start_at, end_at) VALUES (?,?,?,?,?)',
                     (tid, wid, phase or '—', start, end))
    old, new = r['_old'], r['_new']
    changed = {c for c in set(old) | set(new) if old.get(c, 0) != new.get(c, 0)} if r['action'] == '覆盖' else set(new)
    if r['action'] == '覆盖':
        for ev in conn.execute("SELECT e.*, NULL rule_code FROM economic_event e WHERE well_id=? AND period=? AND origin_type='历史导入' "
                               "AND status='有效'", (wid, period)).fetchall():
            if ev['item_code'] in changed:   # 金额没变的科目保留，变了的先冲销再重新导入
                engine.reverse_event(conn, ev, batch_id)
    item = {x['item_code']: x for x in conn.execute('SELECT * FROM cost_item')}
    note = f"单井经营情况模板导入 · {filename} · {r['sheet']}"
    written = 0
    for code, amt in new.items():
        if code not in changed:
            continue
        written += 1
        it = item[code]
        conn.execute(
            "INSERT INTO economic_event (event_date, period, well_id, team_id, item_code, direction, cost_behavior, owner_dept, "
            "origin_type, amount, batch_id, status, source_note) VALUES (?,?,?,?,?,?,?,?, '历史导入', ?,?, '有效', ?)",
            (end, period, wid, tid, code, it['direction'], it['cost_behavior'], it['owner_dept'], amt, batch_id, note))
    log(conn, '经营结果', f'{wid}/{period}', '模板导入' if r['action'] == '导入' else '模板覆盖导入', user['who'],
        note + (f'；变化 {len(changed)} 个科目' if r['action'] == '覆盖' else ''))
    engine.recalc(conn, [wid], [period], '模板导入', user['who'])
    return written
