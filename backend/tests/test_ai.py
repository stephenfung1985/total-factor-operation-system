"""AI 规则解析（第一期）：模型用替身，校验依据文件 → 草稿 → 现有确认流程。"""
import json
import shutil
import subprocess
from io import BytesIO

import pytest
from fastapi.testclient import TestClient


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setenv('TFOS_DB', str(tmp_path / 'test.db'))
    for k in ('TFOS_AI_BASE_URL', 'TFOS_AI_MODEL', 'TFOS_AI_API_KEY'):
        monkeypatch.delenv(k, raising=False)
    from app.db import init_db
    init_db()
    from app.main import app
    return TestClient(app)


FIN, TEAM = {'X-User-Id': 'u_fin'}, {'X-User-Id': 'u_team'}

CONTRACT = """定向井技术服务合同
第一条 服务内容：乙方为甲方提供定向井技术服务。
第五条 计价标准：定向井服务按服务天数计价，单价 38,000 元/天，以签证单为准。
第六条 复杂情况处理费用据实结算，双方另行协商。
第七条 测井服务按 52,000 元/井次结算。
"""

FAKE = {
    'clauses': [
        {'clause_no': '第五条', 'page': '', 'text': '定向井服务按服务天数计价，单价 38,000 元/天',
         'rules': [{'name': '定向井服务日费', 'item_code': 'C2-3-04', 'fact_type': '外包服务', 'line': 'C2-3-04',
                    'pricing_method': '单价', 'unit_price': 38000, 'unit': '天', 'base_item_code': '',
                    'conditions': [{'dimension': '市场', 'operator': '等于', 'value': '西北'},
                                   {'dimension': '区块', 'operator': '等于', 'value': '不存在的区块'}],
                    'combine_mode': '互斥', 'granularity': '逐日', 'valid_from': '2026-11-01',
                    'evidence': '定向井服务按服务天数计价，单价 38,000 元/天', 'confidence': 0.93,
                    'field_confidence': {'科目': 0.95, '触发事实': 0.9, '计价方式': 0.95, '单价': 0.97, '适用条件': 0.9,
                                         '组合方式': 0.9, '计算粒度': 0.9}}]},
        {'clause_no': '第七条', 'page': '',
         'rules': [{'name': '测井', 'item_code': 'C9-99', 'fact_type': '外包服务', 'pricing_method': '单价', 'unit_price': 52000,
                    'unit': '井次', 'evidence': '测井服务按 52,000 元/井次结算', 'confidence': 0.9},
                   {'name': '编造的条款', 'item_code': 'C2-3-14', 'fact_type': '外包服务', 'pricing_method': '单价', 'unit_price': 1,
                    'unit': '井次', 'evidence': '合同里根本没有这一句', 'confidence': 0.9}]},
    ],
    'manual': [{'clause_no': '第六条', 'text': '复杂情况处理费用据实结算，双方另行协商。', 'reason': '据实结算，写不成公式'}],
}


def upload(client, name='定向井合同.txt', data=CONTRACT.encode(), doc_type='外包服务合同', h=FIN):
    return client.post('/api/ai/docs', files={'file': (name, data)}, data={'doc_type': doc_type, 'doc_no': 'DX-2026-07',
                                                                          'counterparty': '某定向服务公司'}, headers=h)


def test_not_configured_records_failed_task(client):
    assert client.get('/api/ai/status', headers=FIN).json()['enabled'] is False
    doc = upload(client).json()
    r = client.post(f"/api/ai/docs/{doc['doc_id']}/parse", headers=FIN)
    assert r.status_code == 502 and '未配置' in r.json()['detail']
    tasks = client.get('/api/ai/tasks', headers=FIN).json()
    assert tasks[0]['status'] == '失败' and tasks[0]['doc_id'] == doc['doc_id']     # 失败也留痕
    assert client.get('/api/ai/status', headers=TEAM).status_code == 403
    assert upload(client, h=TEAM).status_code == 403


def test_upload_rejects_duplicates_and_unsupported(client):
    assert upload(client).status_code == 200
    assert upload(client).status_code == 409                                          # 同一文件指纹
    assert upload(client, name='x.xls', data=b'abc' * 20).status_code == 415
    assert upload(client, name='x.doc', data=b'abc' * 20).status_code == 422                # 扩展名是 .doc 但不是 Word 文件
    assert upload(client, data=CONTRACT.encode(), doc_type='').status_code == 422


def test_docx_text_includes_tables():
    import docx
    from app import aiparse
    d = docx.Document()
    d.add_paragraph('补充协议 第二条 价格调整')
    t = d.add_table(rows=2, cols=3)
    for i, v in enumerate(['项目', '单位', '单价']):
        t.rows[0].cells[i].text = v
    for i, v in enumerate(['定向井服务', '天', '39000']):
        t.rows[1].cells[i].text = v
    buf = BytesIO()
    d.save(buf)
    text, _ = aiparse.extract_text('补充协议.docx', buf.getvalue())
    assert '价格调整' in text and '定向井服务 | 天 | 39000' in text


@pytest.mark.skipif(not shutil.which('textutil') and not shutil.which('antiword'), reason='需要 textutil 或 antiword')
def test_legacy_doc_text(client, tmp_path):
    src = tmp_path / '补充协议.txt'
    src.write_text('补充协议 第二条 价格调整\n定向井服务单价 39,000 元/天，自 2026-11-01 起执行。\n', encoding='utf-8')
    if shutil.which('textutil'):
        subprocess.run(['textutil', '-convert', 'doc', '-output', str(tmp_path / 'a.doc'), str(src)], check=True)
    else:
        pytest.skip('只有 antiword 时无法生成 .doc 样本')
    r = upload(client, name='补充协议.doc', data=(tmp_path / 'a.doc').read_bytes())
    assert r.status_code == 200, r.text
    assert r.json()['chars'] > 20 and r.json()['file_name'] == '补充协议.doc'


def test_parse_creates_checked_drafts(client, monkeypatch):
    from app import ai
    seen = {}

    def fake(messages):
        seen['messages'] = messages
        return '```json\n' + json.dumps(FAKE, ensure_ascii=False) + '\n```'
    monkeypatch.setattr(ai, 'call_model', fake)
    monkeypatch.setenv('TFOS_AI_BASE_URL', 'http://fake/v1')
    monkeypatch.setenv('TFOS_AI_MODEL', 'fake-model')
    doc = upload(client).json()
    out = client.post(f"/api/ai/docs/{doc['doc_id']}/parse", headers=FIN).json()
    # 系统字典随提示词下发，AI 只能用系统里的编码
    user_msg = seen['messages'][1]['content']
    assert 'C2-3-04' in user_msg and '定向井服务按服务天数计价' in user_msg
    assert len(out['created']) == 1 and len(out['skipped']) == 2 and len(out['manual']) == 1
    reasons = ' '.join(s['reason'] for s in out['skipped'])
    assert '科目 C9-99' in reasons and '找不到' in reasons                            # 科目不对、编造的原文都不采纳
    d = out['created'][0]
    assert '适用条件' in d['low_fields']                                              # 区块取值不在基础信息里：标低
    rule = client.get(f"/api/rules/{d['rule_id']}", headers=FIN).json()
    assert rule['rule']['status'] == '待确认' and rule['rule']['origin'] == 'AI提取' and rule['rule']['unit_price'] == 38000
    assert rule['rule']['ai_confidence'] < 0.6 and rule['clause']['hit'] == '定向井服务按服务天数计价，单价 38,000 元/天'
    assert rule['clause']['no'] == '第五条' and rule['rule']['valid_from'] == '2026-11-01'
    assert {c['dimension'] for c in rule['conditions']} == {'市场', '区块'}
    task = client.get('/api/ai/tasks', headers=FIN).json()[0]
    assert task['status'] == '完成' and task['model'] == 'fake-model' and task['draft_count'] == 1
    assert task['adoption'] == {'pending': 1, 'published': 0, 'modified': 0, 'rejected': 0}
    # 草稿走现有流程：确认 → 回算 → 待发布；AI 不能越过人工确认
    c = client.post(f"/api/rules/{d['rule_id']}/confirm", headers=FIN).json()
    assert c['rule']['status'] in ('待发布', '待确认')
    docs = client.get('/api/ai/docs', headers=FIN).json()
    assert docs[0]['parse_status'] == '已解析' and docs[0]['last_task']['draft_count'] == 1


def test_bad_model_output_fails_gracefully(client, monkeypatch):
    from app import ai
    monkeypatch.setattr(ai, 'call_model', lambda m: '抱歉，我无法处理')
    monkeypatch.setenv('TFOS_AI_BASE_URL', 'http://fake/v1')
    monkeypatch.setenv('TFOS_AI_MODEL', 'fake-model')
    doc = upload(client).json()
    r = client.post(f"/api/ai/docs/{doc['doc_id']}/parse", headers=FIN)
    assert r.status_code == 502 and 'JSON' in r.json()['detail']
    assert client.get('/api/ai/docs', headers=FIN).json()[0]['parse_status'] == '解析失败'


def test_evidence_matching_tolerates_pdf_line_breaks():
    from app import aiparse
    text = '第五条 定向井服务按服务\n天数计价，单价 38,000\n元/天。'
    assert aiparse.find_span(text, '定向井服务按服务天数计价，单价 38,000 元/天') is not None
    assert aiparse.find_span(text, '单价 40,000 元/天') is None


WELL_CONTRACT = """W001 井钻井工程施工合同
第三条 合同价款：本井采用单井固定总价，合同总价 1,280 万元（含税），设计井深 5200 米。
第四条 框架价格：其余井钻井按进尺 1,450 元/米结算。
"""

FAKE_MODES = {'clauses': [
    {'clause_no': '第三条', 'rules': [{'name': 'W001 单井总价', 'item_code': 'R02', 'fact_type': '进尺', 'pricing_method': '包干',
                                      'contract_mode': '单井总价', 'unit_price': 12800000, 'unit': '井',
                                      'conditions': [{'dimension': '井号', 'operator': '等于', 'value': 'W001'}],
                                      'evidence': '本井采用单井固定总价，合同总价 1,280 万元（含税）', 'confidence': 0.95}]},
    {'clause_no': '第四条', 'rules': [{'name': '米费', 'item_code': 'R02', 'fact_type': '日历天', 'pricing_method': '单价',
                                      'contract_mode': '米费', 'unit_price': 1450, 'unit': '米',
                                      'evidence': '其余井钻井按进尺 1,450 元/米结算', 'confidence': 0.95}]},
]}


def test_parse_recognizes_contract_mode(client, monkeypatch):
    from app import ai
    seen = {}

    def fake(messages):
        seen['m'] = messages
        return json.dumps(FAKE_MODES, ensure_ascii=False)
    monkeypatch.setattr(ai, 'call_model', fake)
    monkeypatch.setenv('TFOS_AI_BASE_URL', 'http://fake/v1')
    monkeypatch.setenv('TFOS_AI_MODEL', 'fake-model')
    doc = upload(client, name='W001合同.txt', data=WELL_CONTRACT.encode(), doc_type='甲方合同').json()
    out = client.post(f"/api/ai/docs/{doc['doc_id']}/parse", headers=FIN).json()
    assert '单井总价' in seen['m'][0]['content'] and 'W001' in seen['m'][1]['content']   # 计价模式说明与井号字典随提示词下发
    lump, meter = out['created']
    assert not {'单价', '适用条件', '计价方式'} & set(lump['low_fields'])   # 万元写法的总价能在原文找到；限定了井号
    rule = client.get(f"/api/rules/{lump['rule_id']}", headers=FIN).json()
    fields = {x['label']: x['value'] for x in rule['fields']}
    assert rule['rule']['pricing_method'] == '包干' and rule['rule']['unit_price'] == 12800000
    assert fields['计价方式'] == '包干 · 单井总价' and '合同总价 12,800,000 元' in fields['单价']
    assert rule['params']['contract_mode'] == '单井总价'
    assert '计价方式' in meter['low_fields']                # 米费却写成按日历天触发：标低，确认时核对


# ---------- 填报端 AI 解析 ----------

def test_fill_extract_wellinfo(client, monkeypatch):
    from app import ai
    monkeypatch.setenv('TFOS_AI_BASE_URL', 'http://x')
    monkeypatch.setenv('TFOS_AI_MODEL', 'm')
    monkeypatch.setattr(ai, 'call_model', lambda msgs: json.dumps({
        'fields': {'well_id': 'T-1', 'market': '西北', 'well_type': '火星井', 'design_depth': '4,500米', 'spud_sections': 3, 'period': '2026年11月'},
        'milestones': {'开钻1': '2026-11-02 08:30', '乱写': '2026-11-03', '完钻1': '不是时间'}, 'staff': {'合同化': '12人'}}))
    r = client.post('/api/fill/ai-extract/wellinfo', headers=TEAM, files={'file': ('a.txt', '井号 T-1 西北市场 设计井深4500米'.encode())})
    assert r.status_code == 200, r.text
    j = r.json()
    assert j['fields']['well_id'] == 'T-1' and j['fields']['market'] == '西北' and 'well_type' not in j['fields']
    assert j['fields']['design_depth'] == 4500 and j['fields']['period'] == '2026-11'
    assert j['milestones'] == {'开钻1': '2026-11-02 08:30'} and j['staff'] == {'合同化': 12}
    assert any('井型' in n for n in j['notes'])


def test_fill_extract_disabled(client):
    r = client.post('/api/fill/ai-extract/wellinfo', headers=TEAM, files={'file': ('a.txt', b'hello world 123')})
    assert r.status_code == 503
    assert client.get('/api/fill/ai-status', headers=TEAM).json()['enabled'] is False


def test_upload_markdown(client):
    r = upload(client, name='价目表.md', data=('# 结算价目\n\n| 项目 | 单价 |\n|---|---|\n' + CONTRACT).encode())
    assert r.status_code == 200, r.text
    assert r.json()['file_name'] == '价目表.md' and r.json()['chars'] > 20
