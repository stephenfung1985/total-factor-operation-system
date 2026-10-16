import sqlite3
from io import BytesIO

import openpyxl
import pytest
from fastapi.testclient import TestClient


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setenv('TFOS_DB', str(tmp_path / 'test.db'))
    from app.db import init_db
    init_db()
    from app.main import app
    return TestClient(app)


def H(user):
    return {'X-User-Id': user}


def kpi(client, well, period, user='u_fin'):
    return client.get(f'/api/wells/{well}/result?period={period}', headers=H(user)).json()['kpi']


# ---------- 口径与报表 ----------

def test_w001_august_matches_export_template(client):
    k = kpi(client, 'W001', '2026-08')
    # 《单井经营情况导出模板》示例井：总成本 13,836,225.50，形象利润 -6,688,396.66（模板为未舍入的浮点值）；
    # 系统每个科目取整到元，百余个科目累计的取整差在几元以内
    assert k['total_cost'] == pytest.approx(13836225.50, abs=10)
    assert k['net_profit'] == pytest.approx(-6688396.66, abs=10)
    assert k['gross_margin'] == pytest.approx(k['revenue'] - k['variable'])
    assert k['total_cost'] == pytest.approx(k['fixed'] + k['variable'])


def test_money_rounds_half_up_to_whole_yuan():
    from app.db import money
    assert [money(x) for x in (2.5, -2.5, 2.675, 0.49, -0.49, 1832104.12)] == [3, -3, 3, 0, 0, 1832104]


def test_amounts_are_whole_yuan(client):
    """金额类数据保留整数：汇总、成本树、事项、溯源公式、导出文件里的金额都是整数元。"""
    fin = H('u_fin')
    whole = lambda x: float(x).is_integer()   # noqa: E731
    for period in ('2026-08', '2026-10'):
        ov = client.get(f'/api/overview?period={period}', headers=fin).json()
        assert ov['wells'] and all(whole(w[k]) for w in ov['wells'] for k in ('revenue', 'fixed', 'variable', 'total_cost', 'gross_margin', 'net_profit'))
        res = client.get(f'/api/wells/W001/result?period={period}', headers=fin).json()
        assert all(whole(v) for v in res['kpi'].values())
        assert all(whole(n['amount']) for n in res['tree'])
    ev = client.get('/api/wells/W001/items/C2-3-30/events?period=2026-10', headers=fin).json()   # 比例计提
    assert ev['events'] and all(whole(e['amount']) for e in ev['events']) and whole(ev['total'])
    t = client.get(f"/api/events/{ev['events'][0]['event_id']}/trace", headers=fin).json()
    assert t['formula'].endswith(' 元') and '.' not in t['formula'].split('=')[-1]

    r = client.get('/api/wells/W001/export.xlsx?period=2026-10', headers=fin)
    ws = openpyxl.load_workbook(BytesIO(r.content)).active
    amount_rows = [i for i in range(15, ws.max_row + 1) if isinstance(ws.cell(i, 2).value, (int, float))]
    assert amount_rows and all(whole(ws.cell(i, 2).value) for i in amount_rows)
    assert all(ws.cell(i, 2).number_format == '#,##0_ ' for i in amount_rows)


def test_export_xlsx_follows_template_rows(client):
    r = client.get('/api/wells/W001/export.xlsx?period=2026-08', headers=H('u_fin'))
    assert r.status_code == 200
    ws = openpyxl.load_workbook(BytesIO(r.content)).active
    assert ws.max_row == 174
    labels = {ws.cell(i, 1).value: ws.cell(i, 2).value for i in range(1, 175)}
    assert labels['井号'] == 'W001'
    assert labels['30.科研计提'] == pytest.approx(7147828.84 * 0.02, abs=0.5)


def test_mud_service_drill_down_and_trace(client):
    ev = client.get('/api/wells/W001/items/C2-1-02-01/events?period=2026-08', headers=H('u_fin')).json()
    assert len(ev['events']) == 5
    assert ev['total'] == pytest.approx(1832104.12, abs=3)   # 5 笔各取整到元
    t = client.get(f"/api/events/{ev['events'][0]['event_id']}/trace", headers=H('u_fin')).json()
    assert t['clause'] and '进尺' in t['clause']['hit']
    assert t['src_title'] == '来源事实'


# ---------- 两级权限 ----------

def test_team_sees_only_own_wells(client):
    o = client.get('/api/overview?period=2026-08', headers=H('u_team')).json()
    assert [w['well_id'] for w in o['wells']] == ['W001']
    assert client.get('/api/wells/W002/result?period=2026-08', headers=H('u_team')).status_code == 404


def test_permissions_page_is_management_only(client):
    assert client.get('/api/admin/permissions', headers=H('u_fin')).status_code == 200
    assert client.get('/api/admin/permissions', headers=H('u_team')).status_code == 403
    nav = client.get('/api/me', headers=H('u_fin')).json()['nav']
    assert {'perm', 'review', 'rules'} <= {i['id'] for g in nav for i in g['items']}
    m = client.get('/api/admin/permissions', headers=H('u_fin')).json()
    assert [r['id'] for r in m['roles']] == ['finance', 'team'] and set(m['matrix'][0]['cells']) == {'team', 'finance'}


# ---------- 成本日报 → 提交即生效 → 规则引擎 ----------

def _report(depth_start, depth_end, hours=None, date='2026-10-10'):
    return {'well_id': 'W001', 'report_date': date, 'depth_start': depth_start, 'depth_end': depth_end,
            'hours': hours or {'钻进': 16, '起下钻': 4, '循环': 2, '接单根': 1.5, '其他': 0.5},
            'lines': {'C2-1-04': 4.4, 'C2-1-05-01': 20000}}


def test_blocking_checks_prevent_submit(client):
    rep = client.post('/api/reports', json=_report(4824, 4870, hours={'钻进': 10}), headers=H('u_team')).json()
    assert any(c['level'] == '阻断' for c in rep['checks'])
    r = client.post(f"/api/reports/{rep['report_id']}/submit", headers=H('u_team'))
    assert r.status_code == 409


def test_report_flow_generates_events_and_correction_reverses(client):
    before = kpi(client, 'W001', '2026-10')
    rep = client.post('/api/reports', json=_report(4824, 4874), headers=H('u_team')).json()
    rid = rep['report_id']
    out = client.post(f'/api/reports/{rid}/submit', headers=H('u_team')).json()
    assert out['status'] == '已审核' and out['reviewer'] == '系统校验' and out['batch']['created'] > 0   # 提交即生效
    after = kpi(client, 'W001', '2026-10')
    mud_price = round(1832104.12 / 4411, 4)
    image_price = round(7147828.84 / 4411, 4)
    assert after['revenue'] - before['revenue'] == pytest.approx(50 * image_price, abs=0.5)
    assert after['total_cost'] - before['total_cost'] == pytest.approx(
        50 * mud_price + 13666.6667 + 30 * 400 + 16 * 320 + 46 * 55 + 4.4 * 7650 + 20000 * 0.68 + 50 * image_price * 0.031, abs=5)

    # 审核后更正：进尺 50 → 40，原事项冲销、按新事实重算
    new = client.post(f'/api/reports/{rid}/correct', headers=H('u_team')).json()
    assert new['version'] == 2
    client.put(f"/api/reports/{new['report_id']}", json=_report(4824, 4864), headers=H('u_team'))
    out = client.post(f"/api/reports/{new['report_id']}/submit", headers=H('u_team')).json()
    assert out['batch']['reversed'] > 0
    final = kpi(client, 'W001', '2026-10')
    assert final['revenue'] - before['revenue'] == pytest.approx(40 * image_price, abs=0.5)
    old = client.get(f'/api/reports/{rid}', headers=H('u_team')).json()
    assert old['status'] == '已被更正'


# ---------- 规则维护 ----------

def test_rule_confirm_backtest_publish(client):
    lst = client.get('/api/rules', headers=H('u_fin')).json()
    draft = next(r for r in lst if r['rule_code'] == 'RL-C2-3-04')
    assert draft['status'] == '待确认'
    # 规则由管理端维护，基层队无权
    assert client.post(f"/api/rules/{draft['rule_id']}/confirm", headers=H('u_team')).status_code == 403
    d = client.post(f"/api/rules/{draft['rule_id']}/confirm", headers=H('u_fin')).json()
    assert d['backtest']['result'] == '豁免' and d['rule']['status'] == '待发布'
    # 发布必须登记部室提供的变化依据
    assert client.post(f"/api/rules/{draft['rule_id']}/publish", json={}, headers=H('u_fin')).status_code == 422
    d = client.post(f"/api/rules/{draft['rule_id']}/publish", json={'change_basis': '工程技术部 2026-12 号 定向服务价格通知'},
                    headers=H('u_fin')).json()
    assert d['rule']['status'] == '已发布' and '定向服务价格通知' in d['change_basis']
    assert any('变化依据' in (x['comment'] or '') for x in d['logs'])
    # 发布后日报出现“定向服务”行，且外包服务无签证单为阻断
    rep = client.get('/api/reports/current?well_id=W001&date=2026-10-10', headers=H('u_team')).json()
    ld = next(x for x in rep['line_defs'] if x['key'] == 'C2-3-04')
    assert ld['show'] and ld['priced'] and ld['rules'] == ['RL-C2-3-04'] and ld['needs_doc']
    body = _report(4824, 4874)
    body['lines']['C2-3-04'] = 1
    checks = client.post('/api/reports/check', json=body, headers=H('u_team')).json()
    assert any(c['rule'] == '签证单' and c['level'] == '阻断' for c in checks)
    body['notes'] = {'C2-3-04': 'QZ-20261010-01'}
    checks = client.post('/api/reports/check', json=body, headers=H('u_team')).json()
    assert not any(c['level'] == '阻断' for c in checks)


def test_new_version_supersedes_and_recalculates(client):
    lst = client.get('/api/rules', headers=H('u_fin')).json()
    mud = next(r for r in lst if r['rule_code'] == 'RL-C2-1-02-01')
    nv = client.post(f"/api/rules/{mud['rule_id']}/new-version", json={'valid_from': '2026-08-01'}, headers=H('u_fin')).json()
    nid = nv['rule']['rule_id']
    assert nv['rule']['version'] == 2
    client.put(f'/api/rules/{nid}', json={'unit_price': 500}, headers=H('u_fin'))
    d = client.post(f'/api/rules/{nid}/confirm', headers=H('u_fin')).json()
    assert d['backtest']['result'] == '不通过' and d['rule']['status'] == '待确认'
    # 偏差超阈值不能发布
    assert client.post(f'/api/rules/{nid}/publish', json={'change_basis': '调价通知'}, headers=H('u_fin')).status_code == 409


def test_events_are_append_only(tmp_path, client):
    from app.db import connect
    conn = connect()
    with pytest.raises(sqlite3.IntegrityError):
        conn.execute('UPDATE economic_event SET amount=0 WHERE event_id=1')
    with pytest.raises(sqlite3.IntegrityError):
        conn.execute('DELETE FROM economic_event WHERE event_id=1')
    with pytest.raises(sqlite3.IntegrityError):
        conn.execute("UPDATE economic_rule SET unit_price=1 WHERE status='已发布'")
    conn.close()


def test_report_template_roundtrip(client):
    r = client.get('/api/reports/template.xlsx?well_id=W001&start=2026-10-10&end=2026-10-11', headers=H('u_team'))
    assert r.status_code == 200
    wb = openpyxl.load_workbook(BytesIO(r.content))
    assert wb.sheetnames == ['2026-10-10', '2026-10-11']
    res = client.post('/api/reports/import', files={'file': ('t.xlsx', r.content)}, headers=H('u_team')).json()
    assert all(x['ok'] for x in res)
    rep = client.get(f"/api/reports/{res[0]['report_id']}", headers=H('u_team')).json()
    assert rep['status'] == '草稿'


# ---------- 成本日报报量项目 ----------

def test_report_catalog_covers_cost_tree(client):
    cat = client.get('/api/reports/catalog', headers=H('u_team')).json()
    assert [x['id'] for x in cat['sections']] == ['drill', 'energy', 'equip', 'service', 'logistics']
    collected = {x['item_code'] for x in cat['lines']}
    not_collected = {i['item_code'] for g in cat['not_collected'] for i in g['items']}
    assert len(collected) == 55 and len(not_collected) == 64 and not collected & not_collected
    reasons = {g['reason']: [i['item_code'] for i in g['items']] for g in cat['not_collected']}
    assert {'C1-01-01', 'C1-01-02', 'C2-3-18'} <= set(reasons['按基础信息中的人员数量自动核算'])   # 人工、伙食
    assert 'C1-05' in reasons['系统按施工期每天自动生成队天，不用填']      # 制造费
    assert 'C2-3-26' in reasons['物资系统出库单接口带入']                # 套管


def test_new_report_layers_and_carry_over(client):
    rep = client.get('/api/reports/current?well_id=W001&date=2026-10-10', headers=H('u_team')).json()
    shown = {x['key'] for x in rep['line_defs'] if x['show']}
    # 只显示本井当天有已发布规则计价、参与科目费用计算的项目
    assert shown == {x['key'] for x in rep['line_defs'] if x['priced']}
    assert {'C2-1-04', 'C2-1-05-01', 'C2-3-17-01', 'C2-2-02-01'} <= shown and 'C1-01-01' not in shown
    assert 'C2-3-23-03' not in shown and 'C2-1-03' not in shown   # 压裂、润滑油暂无规则，不让基层填
    # 在用设备沿用上一张已审核日报，能耗不沿用
    assert rep['lines'] == {'C2-2-02-01': 1} and rep['notes'] == {'C2-2-02-01': 'DQ-031'}
    # 人员数量取自基础信息（W001：合同化 18 + 市场化 12 + 第三方 16）
    assert rep['staff'] == {'合同化': 18, '市场化': 12, '第三方': 16}


def test_report_checks_service_and_unpriced(client):
    body = _report(4824, 4874)
    body['hours'] = {'钻进': 14, '起下钻': 4, '固井': 2, '接单根': 1.5, '其他': 2.5}
    body['lines']['C2-3-14'] = 1          # 测井：无已发布规则、当天无测井工况
    body['notes'] = {'C2-3-14': 'CJ-01'}
    checks = client.post('/api/reports/check', json=body, headers=H('u_team')).json()
    msgs = {c['rule']: c for c in checks}
    assert msgs['未匹配规则']['level'] == '提示'
    rules = [c['message'] for c in checks if c['rule'] == '服务与工况']
    assert any('测井' in m for m in rules) and any('固井' in m for m in rules)


def test_legacy_line_keys_still_price(client):
    """旧版日报的行名按现行编码计价；旧的“人员出勤”行不再收，人工费按基础信息人员数量核算。"""
    before = kpi(client, 'W001', '2026-10')['total_cost']
    body = _report(4824, 4874)
    body['lines'] = {'人员出勤': 10}
    rid = client.post('/api/reports', json=body, headers=H('u_team')).json()['report_id']
    assert client.get(f'/api/reports/{rid}', headers=H('u_team')).json()['lines'] == {}   # 旧的人员出勤行被忽略
    client.post(f'/api/reports/{rid}/submit', headers=H('u_team'))
    after = kpi(client, 'W001', '2026-10')['total_cost']
    mud_price = round(1832104.12 / 4411, 4)
    image_price = round(7147828.84 / 4411, 4)
    assert after - before == pytest.approx(50 * mud_price + 13666.6667 + (30 * 400 + 16 * 320) + 46 * 55 + 50 * image_price * 0.031, abs=5)
