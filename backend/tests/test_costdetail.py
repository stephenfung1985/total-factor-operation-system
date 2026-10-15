from io import BytesIO

import openpyxl
import pytest
from fastapi.testclient import TestClient


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setenv('TFOS_DB', str(tmp_path / 'test.db'))
    from app.db import init_db
    init_db()
    from app.main import app
    return TestClient(app)


DEPT = {'X-User-Id': 'u_dept'}
FIN = {'X-User-Id': 'u_fin'}
TEAM = {'X-User-Id': 'u_team'}
DIMS = ['period', 'market', 'block', 'team']


def get(client, h, by='period', **q):
    return client.get('/api/cost-detail', params={'by': by, **q}, headers=h)


def test_dept_scope_and_no_revenue(client):
    opts = client.get('/api/cost-detail/options', headers=DEPT).json()
    codes = {i['code'] for i in opts['items']}
    assert codes and not any(c.startswith('R') for c in codes)          # 部室看不到收入
    fin = client.get('/api/cost-detail/options', headers=FIN).json()
    assert len(fin['items']) > len(opts['items'])                       # 经营总览看全部成本科目
    leaf_depts = {i['dept'] for i in opts['items'] if i['is_leaf']}
    assert all('技术' in d.split('/') for d in leaf_depts)               # 部室只含归口科目


@pytest.mark.parametrize('h,status', [(TEAM, 403)])
def test_other_roles_forbidden(client, h, status):
    assert get(client, h).status_code == status
    assert client.get('/api/cost-detail/export.xlsx', headers=h).status_code == status


@pytest.mark.parametrize('by', DIMS)
def test_every_dimension_reconciles_with_overview(client, by):
    from app import analysis, costdetail
    from app.db import transaction
    m = get(client, DEPT, by).json()
    assert m['label'] == costdetail.DIMS[by] and m['columns']
    assert sum(c['total'] for c in m['columns']) == m['grand_total']
    leaves = [r for r in m['rows'] if r['is_leaf']]
    assert sum(r['total'] for r in leaves) == m['grand_total']
    for c in m['columns']:      # 列合计等于该列各叶子科目之和
        assert sum(r['values'][c['key']] for r in leaves) == c['total']
    with transaction() as conn:   # 与经营总览同源：各井各期部室可见成本合计
        user = dict(conn.execute("SELECT * FROM app_user WHERE user_id='u_dept'").fetchone())
        want = 0
        for r in conn.execute('SELECT DISTINCT well_id, period FROM economic_event').fetchall():
            k = analysis.compute(conn, user, r['well_id'], r['period'])
            want += sum(n['amount'] for n in k['tree'] if n['is_leaf'])
    assert m['grand_total'] == want


def test_parent_rollup_and_selection(client):
    m = get(client, DEPT, 'market').json()
    by_code = {r['code']: r for r in m['rows']}
    parents = [r for r in m['rows'] if not r['is_leaf']]
    assert parents
    p = parents[0]
    kids = [r for r in m['rows'] if r['parent'] == p['code']]
    assert p['total'] == sum(k['total'] for k in kids)
    sel = get(client, DEPT, 'market', items=p['code']).json()
    got = {r['code'] for r in sel['rows']}
    assert p['code'] in got and all(k['code'] in got for k in kids)
    assert got <= set(by_code)                       # 只含所选科目及其下级
    assert sel['grand_total'] == sum(r['total'] for r in sel['rows'] if r['is_leaf'])


def test_unknown_or_hidden_item_rejected(client):
    assert get(client, DEPT, items='NOPE').status_code == 400
    assert get(client, DEPT, items='R02').status_code == 400          # 收入科目对部室不可见


def test_period_range(client):
    full = get(client, DEPT, 'period').json()
    periods = [c['key'] for c in full['columns']]
    assert len(periods) >= 2
    one = get(client, DEPT, 'period', start=periods[0], end=periods[0]).json()
    assert [c['key'] for c in one['columns']] == [periods[0]]
    assert one['grand_total'] == full['columns'][0]['total']
    assert get(client, DEPT, start='2026-10', end='2026-08').status_code == 400
    assert get(client, DEPT, start='2026-13').status_code == 400


def test_export_workbook_and_audit_log(client):
    r = client.get('/api/cost-detail/export.xlsx', headers=DEPT)
    assert r.status_code == 200 and 'spreadsheetml' in r.headers['content-type']
    wb = openpyxl.load_workbook(BytesIO(r.content))
    assert wb.sheetnames == ['按时间', '按时间·米成本', '按时间·台月成本', '按市场', '按市场·米成本', '按市场·台月成本',
                             '按区块', '按区块·米成本', '按区块·台月成本', '按队伍', '按队伍·米成本', '按队伍·台月成本', '事项明细']
    assert openpyxl.load_workbook(BytesIO(client.get('/api/cost-detail/export.xlsx', params={'unit': 'false'}, headers=DEPT).content)
                                  ).sheetnames == ['按时间', '按市场', '按区块', '按队伍', '事项明细']
    m = get(client, DEPT, 'team').json()
    ws = wb['按队伍']
    header = [c.value for c in ws[3]]
    assert header[:4] == ['科目编码', '科目名称', '归口部室', '成本性质'] and header[-1] == '合计'
    assert [c['label'] for c in m['columns']] == header[4:-1]
    # 进尺、钻机台月行与米成本、单位成本表
    assert [ws.cell(5, 2).value, ws.cell(6, 2).value] == ['进尺（米）', '钻机台月']
    assert [ws.cell(5, 5 + i).value for i in range(len(m['columns']))] == [c['footage'] for c in m['columns']]
    pm = wb['按队伍·米成本']
    leaf = next(i for i, r in enumerate(m['rows']) if r['is_leaf'] and r['total'])
    assert pm.cell(7 + leaf, 1).value == m['rows'][leaf]['code']
    assert [pm.cell(7 + leaf, 5 + i).value for i in range(len(m['columns']))] == [m['rows'][leaf]['per_meter'][c['key']] for c in m['columns']]
    detail = wb['事项明细']
    assert detail.max_row - 1 == m['event_count']                       # 明细逐笔，与汇总同源
    assert sum(row[13] for row in detail.iter_rows(min_row=2, values_only=True)) == m['grand_total']
    text = ' '.join(str(c.value) for s in wb for row in s.iter_rows() for c in row if c.value is not None)
    assert '形象收入' not in text and '合同收入' not in text            # 部室导出不含收入

    one = client.get('/api/cost-detail/export.xlsx', params={'by': 'block', 'detail': 'false', 'unit': 'false'}, headers=DEPT)
    assert openpyxl.load_workbook(BytesIO(one.content)).sheetnames == ['按区块']
    assert client.get('/api/cost-detail/export.xlsx', params={'by': 'bogus'}, headers=DEPT).status_code == 400

    from app.db import transaction
    with transaction() as c:
        n = c.execute("SELECT COUNT(*) FROM approval_log WHERE object_type='导出' AND object_id='cost-detail'").fetchone()[0]
    assert n == 3      # 默认导出、unit=false 导出、单维度导出各记一条


@pytest.mark.parametrize('by', DIMS)
def test_footage_rig_months_and_unit_costs(client, by):
    m = get(client, FIN, by).json()
    assert m['footage'] > 0 and m['rig_months'] > 0
    assert m['footage'] == pytest.approx(sum(c['footage'] for c in m['columns']), abs=0.05)
    assert m['per_meter'] == pytest.approx(m['cost_with_footage'] / sum(c['footage'] for c in m['columns']), abs=0.01)
    assert m['cost_with_footage'] <= m['grand_total']
    assert m['per_rig'] == pytest.approx(m['grand_total'] / sum(c['rig_months'] for c in m['columns']), abs=0.01)
    for c in m['columns']:
        if c['footage']:
            assert c['per_meter'] == pytest.approx(c['cost_with_footage'] / c['footage'], abs=0.01)
        else:
            assert c['per_meter'] is None
    leaf = next(r for r in m['rows'] if r['is_leaf'] and r['total'])
    assert leaf['per_meter_total'] is not None and leaf['per_meter_total'] <= leaf['total'] / m['footage'] + 0.01


def test_rig_months_follow_work_period_days(client):
    """W001 在 2026-08 整月在队伍施工：台月 = 施工天数 ÷ 当月天数；进尺取已审核日报进尺。"""
    from app import costdetail
    from app.db import transaction
    with transaction() as c:
        rows = [d for d in costdetail._denoms(c, '2026-08', '2026-08') if d['well_id'] == 'W001']
        foot = c.execute("SELECT COALESCE(SUM(quantity),0) FROM business_fact WHERE well_id='W001' AND fact_type='进尺' "
                         "AND status='已审核' AND substr(fact_date,1,7)='2026-08'").fetchone()[0]
    assert sum(d['rig_months'] for d in rows) == pytest.approx(1.0, abs=0.001)
    assert sum(d['footage'] for d in rows) == pytest.approx(foot)
    # 部室也能看到进尺和台月（分母不涉及金额）
    assert get(client, DEPT, 'team').json()['footage'] > 0


def test_meter_cost_ignores_months_without_footage(client):
    """模板导入的历史月份没有进尺，不能进米成本分子，否则会虚高。"""
    import openpyxl as xl
    wb = xl.load_workbook(BytesIO(client.get('/api/wells/W001/export.xlsx?period=2026-08', headers=FIN).content))
    ws = wb.active
    rows = {ws.cell(i, 1).value: i for i in range(1, 19)}
    ws.cell(rows['井号'], 2).value, ws.cell(rows['月份'], 2).value, ws.cell(rows['队号'], 2).value = 'H900', '07', '70101'
    buf = BytesIO()
    wb.save(buf)
    assert client.post('/api/results/import?commit=true', files={'file': ('h.xlsx', buf.getvalue())}, headers=FIN).status_code == 200
    m = get(client, FIN, 'period', start='2026-07', end='2026-07').json()
    col = m['columns'][0]
    assert col['total'] > 0 and col['footage'] == 0 and col['cost_with_footage'] == 0 and col['per_meter'] is None
    assert col['rig_months'] == pytest.approx(1.0, abs=0.001) and col['per_rig'] == pytest.approx(col['total'], abs=0.01)   # 台月成本照算
