from io import BytesIO

import openpyxl
import pytest
from fastapi.testclient import TestClient


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setenv('TFOS_DB', str(tmp_path / 'test.db'))
    from app.db import init_db
    init_db()
    from app.main import app
    return TestClient(app)


FIN = {'X-User-Id': 'u_fin'}


def _sheet_from_export(client, **head):
    """用 W001 8 月的导出文件改表头，模拟一份历史模板。"""
    r = client.get('/api/wells/W001/export.xlsx?period=2026-08', headers=FIN)
    wb = openpyxl.load_workbook(BytesIO(r.content))
    ws = wb.active
    rows = {ws.cell(i, 1).value: i for i in range(1, 19)}
    for k, v in head.items():
        ws.cell(rows[k], 2).value = v
    return wb


def _bytes(wb):
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _post(client, data, commit):
    return client.post(f'/api/results/import?commit={str(commit).lower()}', files={'file': ('历史.xlsx', data)}, headers=FIN)


def test_role_renamed(client):
    assert client.get('/api/me', headers={'X-User-Id': 'u_team'}).json()['role_name'] == '填报端'


def test_blank_template(client):
    r = client.get('/api/results/template.xlsx', headers=FIN)
    ws = openpyxl.load_workbook(BytesIO(r.content)).active
    assert ws.max_row == 174 and ws.cell(1, 1).value == '井号' and ws.title == '批量导入'
    assert ws.max_column == 11      # A 列项目 + B–K 十个井列
    assert all(ws.cell(i, c).value is None for i in range(1, 175) for c in range(2, 12))
    assert ws.cell(30, 11).number_format == ws.cell(30, 2).number_format
    assert client.get('/api/results/template.xlsx', headers={'X-User-Id': 'u_team'}).status_code == 403


def test_import_new_well_and_remark(client):
    expected = client.get('/api/wells/W001/result?period=2026-08', headers=FIN).json()['kpi']
    data = _bytes(_sheet_from_export(client, 井号='H001', 月份='07', 队号='BH50305', 区块='塔中'))
    pre = _post(client, data, False).json()
    s = pre['sheets'][0]
    assert s['action'] == '导入' and s['new_well'] and s['new_team'] and s['period'] == '2026-07'
    assert s['total_cost'] == pytest.approx(expected['total_cost'], abs=0.05)
    assert client.get('/api/overview?period=2026-07', headers=FIN).json()['wells'] == []      # 预览不写库

    assert _post(client, data, True).json()['count'] == {'导入': 1}
    ov = client.get('/api/overview?period=2026-07', headers=FIN).json()['wells']
    assert [(w['well_id'], w['remark'], w['team']) for w in ov] == [('H001', '历史数据井', 'BH50305 钻井队')]
    k = ov[0]
    assert k['total_cost'] == pytest.approx(expected['total_cost'], abs=0.05)
    assert k['net_profit'] == pytest.approx(expected['net_profit'], abs=0.05)
    # 已导入的计提科目不再由比例规则重复计算
    ev = client.get('/api/wells/H001/items/C2-3-30/events?period=2026-07', headers=FIN).json()['events']
    assert [e['origin'] for e in ev] == ['历史导入']
    info = client.get('/api/wells/H001/result?period=2026-07', headers=FIN).json()['info']
    assert info['remark'] == '历史数据井' and '历史.xlsx' in info['import_sources'][0]


def test_reimport_unchanged_then_overwrite(client):
    wb = _sheet_from_export(client, 井号='H002', 月份='07', 队号='70101')
    data = _bytes(wb)
    _post(client, data, True)
    assert _post(client, data, False).json()['sheets'][0]['action'] == '无变化'
    ws = wb.active
    for i in range(19, ws.max_row + 1):
        if ws.cell(i, 1).value == '4.柴油':
            ws.cell(i, 2).value = 200000
    data2 = _bytes(wb)
    pre = _post(client, data2, False).json()['sheets'][0]
    assert pre['action'] == '覆盖' and any('合计' in w for w in pre['warns'])   # 改了明细未改汇总行
    before = client.get('/api/overview?period=2026-07', headers=FIN).json()['wells'][0]['total_cost']
    _post(client, data2, True)
    after = client.get('/api/overview?period=2026-07', headers=FIN).json()['wells'][0]['total_cost']
    assert after - before == pytest.approx(200000 - 129386.68, abs=0.5)
    ev = client.get('/api/wells/H002/items/C2-1-04/events?period=2026-07', headers=FIN).json()['events']
    assert sorted(e['status'] for e in ev) == ['冲销', '已冲销', '有效']
    # 金额没变的科目保留原事项，不冲销重导
    ev2 = client.get('/api/wells/H002/items/C2-1-05-01/events?period=2026-07', headers=FIN).json()['events']
    assert [e['status'] for e in ev2] == ['有效']
    assert [c['item'] for c in pre['changes']] == ['4.柴油'] or any('柴油' in c['item'] for c in pre['changes'])


def test_import_rounds_amounts_to_whole_yuan(client):
    wb = _sheet_from_export(client, 井号='H003', 月份='07', 队号='70101')
    ws = wb.active
    row = next(i for i in range(19, ws.max_row + 1) if ws.cell(i, 1).value == '4.柴油')
    ws.cell(row, 2).value = 100.5
    data = _bytes(wb)
    pre = _post(client, data, False).json()['sheets'][0]
    assert float(pre['total_cost']).is_integer() and float(pre['revenue']).is_integer()
    _post(client, data, True)
    ev = client.get('/api/wells/H003/items/C2-1-04/events?period=2026-07', headers=FIN).json()['events']
    assert [e['amount'] for e in ev] == [101]
    ov = client.get('/api/overview?period=2026-07', headers=FIN).json()['wells'][0]
    assert ov['total_cost'] == pre['total_cost']      # 预览合计与入库后一致


def test_import_blocked_for_system_period(client):
    data = _bytes(_sheet_from_export(client, 月份='10'))
    s = _post(client, data, False).json()['sheets'][0]
    assert s['action'] == '阻断' and any('规则计算' in b for b in s['blocks'])


def test_import_rejects_wrong_layout(client):
    wb = _sheet_from_export(client, 井号='H003', 月份='07', 队号='70101')
    wb.active.cell(30, 1).value = '随便写的项目'
    s = _post(client, _bytes(wb), False).json()['sheets'][0]
    assert s['action'] == '阻断' and '与模板' in s['blocks'][0]
    bad = openpyxl.Workbook()
    bad.active.append(['项目', '金额'])
    assert _post(client, _bytes(bad), False).status_code == 409


def test_export_team_no(client):
    rows = client.get('/api/wells/W001/export/preview?period=2026-08', headers=FIN).json()
    assert next(r['v'] for r in rows if r['k'] == '队号') == '70101'


def test_benchmark_dimensions(client):
    one = client.get('/api/benchmark?period=2026-08&dims=market', headers=FIN).json()
    assert len(one['groups']) == 1 and one['groups'][0]['well_count'] == 4
    assert one['total']['total_cost'] == pytest.approx(one['groups'][0]['total_cost'])
    blk = client.get('/api/benchmark?period=2026-08&dims=block', headers=FIN).json()
    assert {g['keys']['block'] for g in blk['groups']} == {'A区块', 'B区块'}
    assert sum(g['total_cost'] for g in blk['groups']) == pytest.approx(blk['total']['total_cost'])
    multi = client.get('/api/benchmark?period=2026-08&dims=market,block,team', headers=FIN).json()
    assert len(multi['groups']) == 4 and set(multi['groups'][0]['keys']) == {'market', 'block', 'team'}
    g = next(x for x in multi['groups'] if x['keys']['team'] == '70101 钻井队')
    assert g['footage'] == pytest.approx(4411) and g['cost_per_m'] == pytest.approx(g['total_cost'] / 4411)
    assert g['net_rate'] == pytest.approx(g['net_profit'] / g['revenue'])
    allp = client.get('/api/benchmark?period=all&dims=team', headers=FIN).json()
    assert allp['total']['well_months'] == 5     # 8 月 4 口 + 10 月 W001
    assert client.get('/api/benchmark?period=all&dims=team', headers={'X-User-Id': 'u_team'}).status_code == 403


def test_benchmark_specific_objects(client):
    # 只对比两支队伍：70101 与 70103
    r = client.get('/api/benchmark?period=2026-08&dims=team&team=70101 钻井队&team=70103 钻井队', headers=FIN).json()
    assert r['filtered'] and sorted(g['keys']['team'] for g in r['groups']) == ['70101 钻井队', '70103 钻井队']
    assert r['scope']['well_count'] == 2 and r['total']['well_count'] == 4
    assert r['scope']['total_cost'] == pytest.approx(sum(g['total_cost'] for g in r['groups']))
    # 选项联动：选了 B区块，队伍只剩 B区块的两支；区块自身的选项不受影响
    r = client.get('/api/benchmark?period=2026-08&dims=team&block=B区块', headers=FIN).json()
    assert [o['value'] for o in r['options']['team']] == ['70103 钻井队', '70104 钻井队']
    assert [o['value'] for o in r['options']['block']] == ['A区块', 'B区块']
    assert {g['keys']['team'] for g in r['groups']} == {'70103 钻井队', '70104 钻井队'}
    # 区块对区块
    r = client.get('/api/benchmark?period=2026-08&dims=block&block=A区块&block=B区块', headers=FIN).json()
    a, b = sorted(r['groups'], key=lambda g: g['keys']['block'])
    assert a['well_count'] == 2 and b['well_count'] == 2


def _post_many(client, files, commit):
    return client.post(f'/api/results/import?commit={str(commit).lower()}',
                       files=[('file', (n, d)) for n, d in files], headers=FIN)


def _add_well_column(client, wb, col, **head):
    """在工作簿第一个工作表的 col 列放入一口井（取自 W001 导出，改表头），模拟横向多井的批量模板。"""
    ws = wb.active
    src = _sheet_from_export(client, **head).active
    for r in range(1, ws.max_row + 1):
        ws.cell(r, col).value = src.cell(r, 2).value


def test_batch_import_wells_side_by_side_and_multiple_files(client):
    wb = _sheet_from_export(client, 井号='H101', 月份='06', 队号='70101')
    _add_well_column(client, wb, 3, 井号='H102', 月份='06', 队号='70101')
    _add_well_column(client, wb, 4, 井号='H103', 月份='06', 队号='70101')
    a = _bytes(wb)
    pre = _post_many(client, [('a.xlsx', a)], False).json()
    assert pre['count'] == {'导入': 3}
    assert [x['well_id'] for x in pre['sheets']] == ['H101', 'H102', 'H103']
    assert [x['sheet'] for x in pre['sheets']] == [f'{wb.active.title} · {c}列' for c in 'BCD']
    assert _post_many(client, [('a.xlsx', a)], True).json()['count'] == {'导入': 3}
    ov = {w['well_id']: w['remark'] for w in client.get('/api/overview?period=2026-06', headers=FIN).json()['wells']}
    assert ov == {'H101': '历史数据井', 'H102': '历史数据井', 'H103': '历史数据井'}
    # 再次导入同一文件：全部无变化；只改 H102 的柴油：仅该井“覆盖”且只冲销柴油
    ws = wb.active
    row = next(i for i in range(19, ws.max_row + 1) if ws.cell(i, 1).value == '4.柴油')
    ws.cell(row, 3).value = 99999
    mod = _post_many(client, [('a.xlsx', _bytes(wb))], False).json()
    assert [x['action'] for x in mod['sheets']] == ['无变化', '覆盖', '无变化']
    assert [c['item'] for c in mod['sheets'][1]['changes']] == ['4.柴油'] or any('柴油' in c['item'] for c in mod['sheets'][1]['changes'])
    # 同一井同一月在不同文件重复：阻断
    dup = _post_many(client, [('a.xlsx', a), ('c.xlsx', a)], False).json()
    assert dup['count'] == {'无变化': 3, '阻断': 3} and '重复' in dup['sheets'][3]['blocks'][0]
    # 有数据但没填井号的列：阻断
    ws.cell(1, 5).value = None
    ws.cell(row, 5).value = 123
    bad = _post_many(client, [('b.xlsx', _bytes(wb))], False).json()['sheets'][-1]
    assert bad['action'] == '阻断' and any('井号' in b for b in bad['blocks'])


def test_old_single_column_sheets_still_work(client):
    wb = _sheet_from_export(client, 井号='H301', 月份='06', 队号='70101')
    r = _post_many(client, [('old.xlsx', _bytes(wb))], False).json()['sheets']
    assert len(r) == 1 and r[0]['action'] == '导入' and r[0]['sheet'] == wb.active.title


def test_delete_well_and_restore_by_import(client):
    data = _bytes(_sheet_from_export(client, 井号='H201', 月份='06', 队号='70101'))
    _post(client, data, True)
    before = client.get('/api/overview?period=2026-06', headers=FIN).json()
    assert [w['well_id'] for w in before['wells']] == ['H201']
    assert client.delete('/api/wells/H201', headers={'X-User-Id': 'u_team'}).status_code == 403
    assert client.delete('/api/wells/NOPE', headers=FIN).status_code == 404
    r = client.delete('/api/wells/H201', headers=FIN)
    assert r.status_code == 200 and r.json()['reversed'] > 0
    assert client.delete('/api/wells/H201', headers=FIN).status_code == 409
    assert client.get('/api/overview?period=2026-06', headers=FIN).json()['wells'] == []
    assert 'H201' not in [g['keys'].get('well') for g in client.get('/api/benchmark?period=all&dims=well', headers=FIN).json()['groups']]
    # 事项只冲不改：原事项被冲销，仍可追溯
    from app.db import connect
    c = connect()
    st = {r[0] for r in c.execute("SELECT status FROM economic_event WHERE well_id='H201'")}
    c.close()
    assert st == {'已冲销', '冲销'}
    # 历史数据井再次导入即恢复
    pre = _post(client, data, False).json()['sheets'][0]
    assert pre['action'] == '导入' and any('恢复' in w for w in pre['warns'])
    _post(client, data, True)
    assert [w['well_id'] for w in client.get('/api/overview?period=2026-06', headers=FIN).json()['wells']] == ['H201']


def test_delete_registered_well_blocks_import_until_reregistered(client):
    assert client.delete('/api/wells/W002', headers=FIN).status_code == 200
    ov = client.get('/api/overview?period=2026-08', headers=FIN).json()['wells']
    assert 'W002' not in [w['well_id'] for w in ov]
    s = _post(client, _bytes(_sheet_from_export(client, 井号='W002', 月份='06', 队号='70102')), False).json()['sheets'][0]
    assert s['action'] == '阻断' and any('重新登记' in b for b in s['blocks'])


def test_batch_new_wells_share_new_team_and_same_well_two_months(client):
    """批量导入时同一批里的新井共用一支新队伍、同一口新井有多个月份，都不能重复建档。"""
    wb = _sheet_from_export(client, 井号='H401', 月份='05', 队号='BH99999')
    _add_well_column(client, wb, 3, 井号='H402', 月份='05', 队号='BH99999')     # 另一口井，同一支新队伍
    _add_well_column(client, wb, 4, 井号='H401', 月份='06', 队号='BH99999')     # H401 的下一个月
    data = _bytes(wb)
    pre = _post_many(client, [('b.xlsx', data)], False).json()
    assert pre['count'] == {'导入': 3}
    res = _post_many(client, [('b.xlsx', data)], True)
    assert res.status_code == 200 and res.json()['count'] == {'导入': 3}
    from app.db import connect
    c = connect()
    n_team = c.execute("SELECT COUNT(*) FROM team WHERE team_no='BH99999'").fetchone()[0]
    n_well = c.execute("SELECT COUNT(*) FROM well WHERE well_id IN ('H401','H402')").fetchone()[0]
    c.close()
    assert n_team == 1 and n_well == 2
    assert [w['well_id'] for w in client.get('/api/overview?period=2026-05', headers=FIN).json()['wells']] == ['H401', 'H402']
    assert [w['well_id'] for w in client.get('/api/overview?period=2026-06', headers=FIN).json()['wells']] == ['H401']


def test_overview_year_prev_structure_trend_and_history_wells_are_done(client):
    m8 = client.get('/api/overview?period=2026-08', headers=FIN).json()
    m10 = client.get('/api/overview?period=2026-10', headers=FIN).json()
    year = client.get('/api/overview?period=2026', headers=FIN).json()
    assert year['months'] == ['2026-08', '2026-10']
    assert year['sum']['revenue'] == m8['sum']['revenue'] + m10['sum']['revenue']
    w1 = next(w for w in year['wells'] if w['well_id'] == 'W001')
    assert w1['last_period'] == '2026-10'
    # 环比：十月对八月；八月之前没有数据
    assert m10['prev']['period'] == '2026-08' and m10['prev']['sum'] == m8['sum'] and m8['prev'] is None
    # 成本科目结构：二级科目合计等于总成本
    assert sum(d['amount'] for d in m8['structure']) == m8['sum']['total_cost']
    tr = client.get('/api/overview/trend', headers=FIN).json()
    assert [t['period'] for t in tr] == ['2026-08', '2026-10'] and tr[0]['revenue'] == m8['sum']['revenue']
    # 历史数据井默认为完成井
    _post(client, _bytes(_sheet_from_export(client, 井号='H501', 月份='07', 队号='70101')), True)
    h = next(w for w in client.get('/api/overview?period=2026-07', headers=FIN).json()['wells'] if w['well_id'] == 'H501')
    assert h['remark'] == '历史数据井' and h['status'] == '完井'


def test_benchmark_multiple_periods(client):
    one = client.get('/api/benchmark?period=2026-08&dims=team', headers=FIN).json()
    both = client.get('/api/benchmark?period=2026-08&period=2026-10&dims=team', headers=FIN).json()
    allp = client.get('/api/benchmark?period=all&dims=team', headers=FIN).json()
    assert one['selected'] == ['2026-08'] and both['selected'] == ['2026-08', '2026-10'] and both['period'] == '2026-08,2026-10'
    assert both['scope']['well_months'] == allp['scope']['well_months'] > one['scope']['well_months']
    assert both['scope']['revenue'] == allp['scope']['revenue']
    # 整年与默认（不传）等价于全部
    year = client.get('/api/benchmark?period=2026&dims=team', headers=FIN).json()
    assert year['selected'] == ['2026-08', '2026-10'] and year['scope']['revenue'] == allp['scope']['revenue']
    assert client.get('/api/benchmark?dims=team', headers=FIN).json()['selected'] == allp['selected']
    assert [o['value'] for o in both['period_options']] == ['2026-08', '2026-10'] and both['period_options'][0]['wells'] == 4
    # 不存在的期间：没有对象，不报错
    assert client.get('/api/benchmark?period=1999-01&dims=team', headers=FIN).json()['groups'] == []


def test_benchmark_rank_metrics(client):
    r = client.get('/api/benchmark?period=2026-08&dims=team', headers=FIN).json()
    gs = r['groups']
    for g in gs:   # 口径：毛益额 = 收入 − 变动成本，毛益率 = 毛益额 ÷ 收入，净利润 = 收入 − 固定 − 变动
        assert g['gross_margin'] == pytest.approx(g['revenue'] - g['variable'])
        assert g['net_profit'] == pytest.approx(g['revenue'] - g['fixed'] - g['variable'])
        assert g['gm_rate'] == pytest.approx(g['gross_margin'] / g['revenue'])
        assert g['footage'] == pytest.approx(sum(i['footage'] for i in g['items']))
    for m in ('net_rate', 'footage', 'revenue', 'gross_margin', 'gm_rate', 'net_profit'):
        vals = [gs[i][m] for i in r['rank'][m]]
        assert sorted(r['rank'][m]) == list(range(len(gs))) and vals == sorted(vals, reverse=True)
    vals = [gs[i]['cost_per_m'] for i in r['rank']['cost_per_m']]
    assert vals == sorted(vals)
    assert r['total']['group_count'] == 4 and r['scope']['group_count'] == 4
    f = client.get('/api/benchmark?period=2026-08&dims=team&block=B区块', headers=FIN).json()
    assert f['scope']['group_count'] == 2 and f['total']['group_count'] == 4


def test_ranking_skips_missing_values():
    from app.analysis import ranking
    gs = [{'gm_rate': 0.1, 'cost_per_m': None}, {'gm_rate': None, 'cost_per_m': 300}, {'gm_rate': 0.3, 'cost_per_m': 200}]
    assert ranking(gs, 'gm_rate') == [2, 0]
    assert ranking(gs, 'cost_per_m') == [2, 1]
