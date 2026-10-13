"""从《单井经营情况导出模板》和设计文档附录 A 生成成本项目树（JSON）。

用法：python3 tools/build_cost_tree.py 单井经营情况导出模板.xlsx "钻井单井经营管控系统 核心数据库表设计.md" > tree.json

按 A 列的项目名称识别层级，不依赖固定行号：
  一、固定成本 / 二、变动成本         → 第 1 级
  1.xxx（固定）/（一）xxx（变动）      → 第 2 级
  （1）xxx（固定）/ 1.xxx（变动）      → 第 3 级
  （1）xxx（变动）                     → 第 4 级
“使用权资产租赁”及其子项即使出现也会跳过（设计文档已取消该项）。
"""
import json
import re
import sys

import openpyxl

CN_SECTION = re.compile(r'^（[一二三四五六七八九十]+）')
NUM_ITEM = re.compile(r'^(\d+)\.(.*)$')
SUB_ITEM = re.compile(r'^（\d+）')


def load_appendix(design_path):
    appx = {}
    for line in open(design_path, encoding='utf-8'):
        m = re.match(r'^\| (C[\d-]+) \| ([^|]+) \| ([A-Z]+) \| [^|]* \| ([^|]+) \| ([^|]+) \| ([^|]+) \|', line)
        if m:
            appx[m.group(1)] = dict(col=m.group(3), t=m.group(4).strip(), d=m.group(6).strip())
    return appx


def main(tpl_path, design_path):
    ws = openpyxl.load_workbook(tpl_path).active
    rows = [(str(ws.cell(r, 1).value or '').strip(), ws.cell(r, 2).value) for r in range(1, ws.max_row + 1)]
    labels = [name for name, _ in rows]
    start, end = labels.index('一、固定成本'), labels.index('经营数据')
    head = dict(rows[:start])

    nodes = []
    behavior = None
    fixed_no = 0
    l2 = l3 = None
    sub = 0
    sec_no = 0
    skipping = False
    for name, val in rows[start:end]:
        amt = float(val or 0)
        if name == '一、固定成本':
            behavior = 'F'
            nodes.append(dict(c='C1', n=name, lv=1, b='F', p=None))
            continue
        if name == '二、变动成本':
            behavior = 'V'
            skipping = False
            nodes.append(dict(c='C2', n=name, lv=1, b='V', p=None))
            continue
        if behavior == 'F':
            m = NUM_ITEM.match(name)
            if m:
                skipping = '使用权资产租赁' in name
                if skipping:
                    continue
                fixed_no += 1
                l2, sub = 'C1-%02d' % fixed_no, 0
                nodes.append(dict(c=l2, n='%d.%s' % (fixed_no, m.group(2)), lv=2, b='F', p='C1', a=amt))
            elif SUB_ITEM.match(name) and not skipping:
                sub += 1
                nodes.append(dict(c='%s-%02d' % (l2, sub), n=name, lv=3, b='F', p=l2, a=amt))
        else:
            if CN_SECTION.match(name):
                sec_no += 1
                l2 = 'C2-%d' % sec_no
                nodes.append(dict(c=l2, n=name, lv=2, b='V', p='C2'))
                continue
            m = NUM_ITEM.match(name)
            if m:
                l3, sub = '%s-%02d' % (l2, int(m.group(1))), 0
                nodes.append(dict(c=l3, n=name, lv=3, b='V', p=l2, a=amt))
            elif SUB_ITEM.match(name):
                sub += 1
                nodes.append(dict(c='%s-%02d' % (l3, sub), n=name, lv=4, b='V', p=l3, a=amt))

    appx = load_appendix(design_path)
    parents = {n['p'] for n in nodes if n['p']}
    out = []
    for n in nodes:
        n['leaf'] = n['c'] not in parents
        if not n['leaf']:
            n.pop('a', None)
        key = n['c'] if n['c'] in appx else (n['p'] if n['p'] in appx else None)
        if key:
            n.update(d=appx[key]['d'], t=appx[key]['t'], col=appx[key]['col'])
        if 'a' in n:
            n['a'] = round(n['a'], 2)
        out.append({k: v for k, v in n.items() if v is not None})

    json.dump(dict(tree=out, rev={'contract': head.get('合同收入（元）'), 'image': head.get('形象收入（元）')}),
              sys.stdout, ensure_ascii=False, separators=(',', ':'))
    print('%d nodes, %d leaves' % (len(out), sum(1 for n in out if n['leaf'])), file=sys.stderr)


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
